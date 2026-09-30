"""台本テキスト → スライド＋AI音声＋テロップ＋BGM入りのMP4 を作る。

使い方: python3 build.py scripts/Q1第1回.txt Q1第1回 "第1回　エアコン清掃独立のリアル"
（音声入り .pptx を作る部品も残してあるが、今は MP4 だけを出している）
"""
import copy
import json
import os
import re
import subprocess
import sys
import wave

import numpy as np
import imageio_ffmpeg
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.media import Video
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
SR = 48000
FONT = "Noto Sans CJK JP"
HERE = os.path.dirname(os.path.abspath(__file__))
LOGO_WHITE = os.path.join(HERE, "logo_white.png")
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xE8, 0x83, 0x3A)
INK = RGBColor(0x22, 0x2B, 0x36)
PAPER = RGBColor(0xFA, 0xF7, 0xF2)
PLACEHOLDER = re.compile(r"\[三浦が(確認|記入)[^\]]*\]")

LEAD_IN = 0.6      # スライドが出てから話し始めるまで（秒）
LINE_GAP = 0.25    # 行と行のあいだ
PARA_GAP = 0.6     # 段落と段落のあいだ
TAIL = 1.0         # 話し終わってから次へ進むまで


# ---------- 台本の読み込み ----------
def parse(path):
    slides = []
    for block in open(path, encoding="utf-8").read().split("=== ")[1:]:
        title, rest = block.split("\n", 1)
        screen, voice = rest.split("[voice]")
        screen = [l for l in screen.replace("[screen]", "").strip().splitlines() if l.strip()]
        paras = [[l.strip() for l in p.splitlines() if l.strip()]
                 for p in voice.strip().split("\n\n")]
        slides.append({"title": title.strip(), "screen": screen, "paras": [p for p in paras if p]})
    return slides


# ---------- 音声 ----------
# 読み間違いの直し（テロップはそのまま、読み上げだけ差し替える）
READING_FIX = {
    "資格はいりません": "資格は要りません",  # 「はいり（入り）」と読まれるのを防ぐ
    "たくさんの方": "たくさんのかた",        # 人を指す「方」は「かた」
    "くれた方": "くれたかた",
    "オーナーの方": "オーナーのかた",
    "他責": "たせき",                    # 「たせめ」と読まれる
    "〇〇": "まるまる",                  # 「ゼロゼロ」と読まれる
    "〇万円": "まるまんえん",            # 「レーマンエン」と読まれる
}


def reading_report(slides):
    """読み間違いやすい語（方・は）がどう読まれるかを一覧にする。"""
    import pyopenjtalk
    for i, s in enumerate(slides, 1):
        for para in s["paras"]:
            for line in para:
                if PLACEHOLDER.search(line):
                    continue
                text = line
                for wrong, right in READING_FIX.items():
                    text = text.replace(wrong, right)
                for w in pyopenjtalk.run_frontend(text):
                    if "方" in w["string"] or ("は" in w["string"] and "ハ" in w["pron"]):
                        print(f"[読み確認] スライド{i}: {w['string']}={w['pron']} | {line}")


def tts(text):
    import pyopenjtalk
    for wrong, right in READING_FIX.items():
        text = text.replace(wrong, right)
    x, sr = pyopenjtalk.tts(text, speed=1.08)
    x = x.astype(np.float32) / 32768.0
    # 48kHz のまま使う（pyopenjtalk は 48kHz）
    assert sr == SR
    return x


def make_bgm(seconds):
    """やわらかいパッド音のBGM（自作・著作権フリー）。C - Am - F - G を8秒ずつ。"""
    chords = [[261.6, 329.6, 392.0], [220.0, 261.6, 329.6],
              [174.6, 220.0, 261.6], [196.0, 246.9, 293.7]]
    bar = 8.0
    n = int(seconds * SR)
    t = np.arange(n) / SR
    out = np.zeros(n, dtype=np.float32)
    for i in range(int(np.ceil(seconds / bar)) + 1):
        start = i * bar - 1.0
        chord = chords[i % 4]
        s0, s1 = max(0, int(start * SR)), min(n, int((start + bar + 2.0) * SR))
        if s0 >= s1:
            continue
        tt = t[s0:s1] - start
        env = np.clip(tt / 2.0, 0, 1) * np.clip((bar + 2.0 - tt) / 2.0, 0, 1)
        for f in chord + [chord[0] / 2]:
            out[s0:s1] += (np.sin(2 * np.pi * f * t[s0:s1])
                           + 0.25 * np.sin(2 * np.pi * 2 * f * t[s0:s1])) * env
    # ゆっくり揺らす
    out *= 0.85 + 0.15 * np.sin(2 * np.pi * 0.1 * t)
    return out / np.max(np.abs(out))


def write_m4a(x, path):
    wav = path + ".wav"
    with wave.open(wav, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(SR)
        w.writeframes((np.clip(x, -1, 1) * 32767).astype(np.int16).tobytes())
    subprocess.run([FFMPEG, "-y", "-loglevel", "error", "-i", wav, "-c:a", "aac",
                    "-b:a", "128k", path], check=True)
    return wav


def build_audio(slides, outdir, bgm_level=0.06):
    """スライドごとに『ナレーション＋BGM』の音声を作り、テロップのタイミングを記録する。"""
    for s in slides:
        chunks, captions, t = [np.zeros(int(LEAD_IN * SR), np.float32)], [], LEAD_IN
        for pi, para in enumerate(s["paras"]):
            for li, line in enumerate(para):
                if PLACEHOLDER.search(line):
                    continue  # 三浦さんの記入待ちは読まない
                x = tts(line)
                d = len(x) / SR
                captions.append({"text": line, "start": t, "end": t + d})
                chunks.append(x)
                t += d
                gap = LINE_GAP if li < len(para) - 1 else PARA_GAP
                chunks.append(np.zeros(int(gap * SR), np.float32))
                t += gap
        voice = np.concatenate(chunks)
        voice = np.concatenate([voice, np.zeros(int(TAIL * SR), np.float32)])
        voice = 0.9 * voice / max(1e-6, np.max(np.abs(voice)))
        s["voice"], s["captions"], s["dur"] = voice, captions, len(voice) / SR

    total = sum(s["dur"] for s in slides)
    bgm = make_bgm(total + 1) * bgm_level
    pos = 0
    for i, s in enumerate(slides):
        n = len(s["voice"])
        mix = s["voice"] + bgm[pos:pos + n]
        # つなぎ目のプチ音を防ぐ短いフェード
        f = int(0.05 * SR)
        mix[:f] *= np.linspace(0, 1, f)
        mix[-f:] *= np.linspace(1, 0, f)
        pos += n
        s["audio"] = os.path.join(outdir, f"slide{i + 1:02d}.m4a")
        s["wav"] = write_m4a(mix, s["audio"])


# ---------- スライド ----------
def set_font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def textbox(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, color, bold, align) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * 0.55)
        # 記入待ちの部分は黄色いマーカーで目立たせる
        pos = 0
        for m in PLACEHOLDER.finditer(text):
            if m.start() > pos:
                set_font(p.add_run(), size, color, bold)
                p.runs[-1].text = text[pos:m.start()]
            r = p.add_run()
            r.text = m.group(0)
            set_font(r, size * 0.8, RGBColor(0x7A, 0x4B, 0x00), True)
            hl = etree.SubElement(r._r.get_or_add_rPr(), qn("a:highlight"))
            etree.SubElement(hl, qn("a:srgbClr")).set("val", "FFE066")
            pos = m.end()
        if pos < len(text):
            r = p.add_run()
            r.text = text[pos:]
            set_font(r, size, color, bold)
    return tb


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, alpha=None):
    sp = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    if alpha is not None:
        clr = sp.fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
    return sp


def draw_slide(prs, s, idx, lesson, with_captions=True):
    W, H = prs.slide_width, prs.slide_height
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    if idx == 0:
        bg.fore_color.rgb = NAVY
        rect(slide, int(W * 0.08), int(H * 0.30), int(W * 0.08), Emu(Pt(6)), ORANGE)
        lines = s["screen"]
        textbox(slide, int(W * 0.08), int(H * 0.33), int(W * 0.84), int(H * 0.40), [
            (lines[0], 26, RGBColor(0xC9, 0xD6, 0xE8), False, PP_ALIGN.LEFT),
            (lines[1], 48, RGBColor(0xFF, 0xFF, 0xFF), True, PP_ALIGN.LEFT),
            (lines[2], 26, RGBColor(0xFF, 0xC8, 0x9A), False, PP_ALIGN.LEFT),
        ])
        slide.shapes.add_picture(LOGO_WHITE, int(W * 0.08), int(H * 0.08), height=int(H * 0.10))
    else:
        bg.fore_color.rgb = PAPER
        rect(slide, 0, 0, W, int(H * 0.17), NAVY)
        rect(slide, 0, int(H * 0.17), W, Emu(Pt(4)), ORANGE)
        textbox(slide, int(W * 0.05), 0, int(W * 0.60), int(H * 0.17),
                [(s["title"], 30, RGBColor(0xFF, 0xFF, 0xFF), True, PP_ALIGN.LEFT)],
                anchor=MSO_ANCHOR.MIDDLE)
        logo_h = int(H * 0.058)
        pic = slide.shapes.add_picture(LOGO_WHITE, 0, int(H * 0.025), height=logo_h)
        pic.left = int(W * 0.965) - pic.width
        textbox(slide, int(W * 0.62), int(H * 0.09), int(W * 0.345), int(H * 0.07),
                [(f"{lesson}　{idx + 1} / {s['total']}", 12, RGBColor(0xC9, 0xD6, 0xE8), False, PP_ALIGN.RIGHT)])
        body = s["screen"]
        n = sum(not l.startswith("※") for l in body) + sum(len(l) > 32 for l in body)
        size = 30 if n <= 4 else 26 if n <= 6 else 22
        out = []
        for l in body:
            if l.startswith("・"):
                out.append(("•  " + l[1:], size * (0.88 if len(l) > 28 else 1), INK, False, PP_ALIGN.LEFT))
            elif l.startswith("※"):
                out.append((l, size * 0.7, RGBColor(0x5C, 0x66, 0x73), False, PP_ALIGN.LEFT))
            elif l.startswith("→"):
                out.append((l, size, ORANGE, True, PP_ALIGN.LEFT))
            else:
                out.append((l, size + 4, NAVY, True, PP_ALIGN.LEFT))
        rect(slide, int(W * 0.05), int(H * 0.25), Emu(Pt(5)), int(H * 0.50), ORANGE)
        textbox(slide, int(W * 0.075), int(H * 0.23), int(W * 0.87), int(H * 0.56), out)

    caption_ids = []
    if with_captions:
        for c in s["captions"]:
            box = rect(slide, int(W * 0.06), int(H * 0.815), int(W * 0.88), int(H * 0.13),
                       RGBColor(0x10, 0x14, 0x1A), MSO_SHAPE.ROUNDED_RECTANGLE, alpha=0.78)
            box.adjustments[0] = 0.18
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = c["text"]
            set_font(r, 22 if len(c["text"]) <= 34 else 19, RGBColor(0xFF, 0xFF, 0xFF), True)
            box.name = "テロップ"
            caption_ids.append(box.shape_id)
    return slide, caption_ids


# ---------- 音声の埋め込みとタイミング ----------
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NSMAP = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
         'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')


def embed_audio(slide, audio_path, icon_path, W):
    media = Video.from_path_or_file_like(audio_path, "audio/mp4")
    media_part = slide.part.package.get_or_add_media_part(media)
    r_link = slide.part.relate_to(media_part, RT.AUDIO)
    r_media = slide.part.relate_to(media_part, RT.MEDIA)
    _, r_img = slide.part.get_or_add_image_part(icon_path)
    sid = max(sp.shape_id for sp in slide.shapes) + 1
    size = 457200
    xml = f'''<p:pic {NSMAP}>
  <p:nvPicPr>
    <p:cNvPr id="{sid}" name="ナレーション"><a:hlinkClick r:id="" action="ppaction://media"/></p:cNvPr>
    <p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>
    <p:nvPr>
      <a:audioFile r:link="{r_link}"/>
      <p:extLst><p:ext uri="{{DAA4B4D4-6D71-4841-9C94-3DA1FCC6ED3A}}">
        <p14:media xmlns:p14="{P14}" r:embed="{r_media}"/></p:ext></p:extLst>
    </p:nvPr>
  </p:nvPicPr>
  <p:blipFill><a:blip r:embed="{r_img}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
  <p:spPr><a:xfrm><a:off x="{W + 200000}" y="200000"/><a:ext cx="{size}" cy="{size}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
</p:pic>'''
    slide.shapes._spTree.append(etree.fromstring(xml))
    return sid


def add_timing(slide, audio_id, dur, captions, caption_ids):
    ids = iter(range(3, 10000))
    ms = lambda sec: str(int(round(sec * 1000)))

    def effect(cls, spid, delay, visibility, node="withEffect"):
        a, b = next(ids), next(ids)
        return (f'<p:par><p:cTn id="{a}" presetID="1" presetClass="{cls}" presetSubtype="0" '
                f'fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
                f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst>'
                f'<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="{visibility}"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')

    outer, inner = next(ids), next(ids)
    a, b = next(ids), next(ids)
    play = (f'<p:par><p:cTn id="{a}" presetID="1" presetClass="mediacall" presetSubtype="0" '
            f'fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'<p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr>'
            f'<p:cTn id="{b}" dur="{ms(dur)}" fill="hold"/><p:tgtEl><p:spTgt spid="{audio_id}"/>'
            f'</p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>')
    fx = [play]
    for c, spid in zip(captions, caption_ids):
        fx.append(effect("entr", spid, ms(c["start"]), "visible"))
        fx.append(effect("exit", spid, ms(c["end"] + 0.15), "hidden"))
    media_node = next(ids)
    xml = f'''<p:timing {NSMAP}><p:tnLst><p:par>
  <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
    <p:seq concurrent="1" nextAc="seek">
      <p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
        <p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>
          <p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
          <p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>
            <p:childTnLst>{"".join(fx)}</p:childTnLst></p:cTn></p:par>
        </p:childTnLst></p:cTn></p:par>
      </p:childTnLst></p:cTn>
      <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
      <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
    </p:seq>
    <p:audio><p:cMediaNode vol="80000"><p:cTn id="{media_node}" fill="hold" display="0">
      <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
      <p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst>
    </p:cTn><p:tgtEl><p:spTgt spid="{audio_id}"/></p:tgtEl></p:cMediaNode></p:audio>
  </p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    sld = slide._element
    trans = etree.fromstring(f'<p:transition {NSMAP} spd="med" advTm="{ms(dur)}"><p:fade/></p:transition>')
    anchor = sld.find(qn("p:clrMapOvr"))
    anchor.addnext(trans)
    trans.addnext(etree.fromstring(xml))


def speaker_icon(path):
    im = Image.new("RGBA", (96, 96), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.polygon([(14, 36), (34, 36), (58, 14), (58, 82), (34, 60), (14, 60)], fill=(31, 58, 95, 255))
    d.arc((52, 28, 84, 68), -60, 60, fill=(31, 58, 95, 255), width=6)
    im.save(path)


def new_prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    return prs


# ---------- 確認用MP4 ----------
def render_mp4(slides, lesson, outdir, mp4_path):
    """テロップなしのスライドを画像にして、テロップを重ね、音声とつないでMP4にする。"""
    prs = new_prs()
    for i, s in enumerate(slides):
        draw_slide(prs, s, i, lesson, with_captions=False)
    base = os.path.join(outdir, "base.pptx")
    prs.save(base)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, base],
                   check=True, capture_output=True)
    import pymupdf as fitz
    for i, page in enumerate(fitz.open(os.path.join(outdir, "base.pdf"))):
        pix = page.get_pixmap(matrix=fitz.Matrix(1920 / page.rect.width, 1080 / page.rect.height))
        pix.save(os.path.join(outdir, f"img-{i + 1:02d}.png"))
    font = ImageFont.truetype("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc", 40, index=0)
    segs, n = [], 0
    for i, s in enumerate(slides):
        img = Image.open(os.path.join(outdir, f"img-{i + 1:02d}.png")).convert("RGB")
        marks = [0.0] + sorted({x for c in s["captions"] for x in (c["start"], c["end"] + 0.15)}) + [s["dur"]]
        for a, b in zip(marks, marks[1:]):
            if b - a < 0.02:
                continue
            frame = img.copy()
            cur = [c for c in s["captions"] if c["start"] <= a + 0.01 < c["end"] + 0.15]
            if cur:
                ov = Image.new("RGBA", frame.size, (0, 0, 0, 0))
                d = ImageDraw.Draw(ov)
                d.rounded_rectangle((115, 880, 1805, 1020), 24, fill=(16, 20, 26, 200))
                text = cur[0]["text"]
                w = d.textlength(text, font=font)
                d.text(((1920 - w) / 2, 928), text, font=font, fill="white")
                frame = Image.alpha_composite(frame.convert("RGBA"), ov).convert("RGB")
            p = os.path.join(outdir, f"f{n:04d}.png")
            frame.save(p)
            segs.append((p, b - a))
            n += 1
    lst = os.path.join(outdir, "frames.txt")
    with open(lst, "w") as f:
        for p, d in segs:
            f.write(f"file '{p}'\nduration {d:.3f}\n")
        f.write(f"file '{segs[-1][0]}'\n")
    alst = os.path.join(outdir, "audio.txt")
    with open(alst, "w") as f:
        for s in slides:
            f.write(f"file '{s['wav']}'\n")
    subprocess.run([FFMPEG, "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", lst,
                    "-f", "concat", "-safe", "0", "-i", alst, "-vf", "fps=30,format=yuv420p",
                    "-c:v", "libx264", "-preset", "medium", "-crf", "20", "-c:a", "aac", "-b:a", "160k",
                    "-shortest", mp4_path], check=True)


def main():
    src, name, lesson = sys.argv[1], sys.argv[2], sys.argv[3]
    outdir = os.path.abspath(name + "_work")
    os.makedirs(outdir, exist_ok=True)
    slides = parse(src)
    for s in slides:
        s["total"] = len(slides)
    reading_report(slides)
    build_audio(slides, outdir)

    render_mp4(slides, lesson, outdir, name + ".mp4")
    json.dump([{"title": s["title"], "dur": round(s["dur"], 1),
                "captions": [(c["text"], round(c["start"], 2), round(c["end"], 2)) for c in s["captions"]]}
               for s in slides], open(os.path.join(outdir, "timing.json"), "w"), ensure_ascii=False, indent=1)
    print("total", round(sum(s["dur"] for s in slides) / 60, 1), "min")


if __name__ == "__main__":
    main()
